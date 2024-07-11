pip install opencv-python-headless
pip install pandas
pip install paddlepaddle
pip install paddleocr
pip install openpyxl


import cv2
import pandas as pd
from paddleocr import PPStructure
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage

# Initialize PPStructure for table extraction with recovery and OCR results
table_engine = PPStructure(recovery=True, return_ocr_result_in_table=True)

# Create and save an Excel workbook to store the results
output = '/content/output.xlsx'
Workbook().save(output)
book = load_workbook(output)
writer = pd.ExcelWriter(output, engine='openpyxl')
writer.book = book

# Process images in a loop
for n in range(1, 5):
    print('image', n)
    img_path = f'/content/{n} (1).png'
    img = cv2.imread(img_path)
    result = table_engine(img)

    # Create an image object for openpyxl
    xlimg = XLImage(img_path)

    i = 1
    for line in result:
        # Remove the 'img' key from the result
        line.pop('img')
        # Check if the line is a table
        if line.get("type") == "table":
            # Extract HTML table and convert to DataFrame
            html_table = line.get("res").get("html")
            html_data = pd.read_html(html_table)
            df = pd.DataFrame(html_data[0])

            # Write DataFrame to Excel and add the image to the sheet
            df.to_excel(writer, sheet_name=f"image {n} table {i}", index=1)
            book[f"image {n} table {i}"].add_image(xlimg, 'A100')
            i += 1

# Save the Excel workbook
writer.save()